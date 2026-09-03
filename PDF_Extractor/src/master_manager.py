"""
Master Manager: Maintains and formats master/PAIMANA_Master.xlsx according to the exact 55-column schema.
"""
import os
import yaml
import logging
from typing import List, Dict, Any, Optional, Set, Tuple
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

logger = logging.getLogger(__name__)

DEFAULT_MASTER_PATH = "master/PAIMANA_Master.xlsx"
SHEET_NAME = "Project_Monthly_Data"

EXACT_55_COLUMNS = [
    "project_id",
    "project_key",
    "project_name",
    "agency",
    "ministry_department",
    "sector",
    "state",
    "legacy_ocms_code",
    "pmgid",
    "report_month",
    "approval_start",
    "original_target_doc",
    "revised_doc",
    "project_age_months",
    "original_duration_months",
    "planned_remaining_months",
    "revised_remaining_months",
    "schedule_extension_months",
    "extension_rate_pct",
    "original_cost_crore",
    "revised_cost_crore",
    "cumulative_expenditure_crore",
    "cost_overrun_pct",
    "expenditure_ratio_pct",
    "cost_escalation_crore",
    "cost_escalation_ratio",
    "remaining_budget_crore",
    "expenditure_velocity_crore_month",
    "physical_progress_pct",
    "remaining_work_pct",
    "physical_progress_delta_1m",
    "physical_progress_delta_3m",
    "physical_progress_delta_6m",
    "progress_velocity_3m",
    "progress_velocity_6m",
    "progress_trend_slope",
    "physical_financial_gap_pct",
    "physical_to_expenditure_ratio",
    "progress_expenditure_mismatch_flag",
    "high_expenditure_low_progress_flag",
    "days_to_original_target",
    "days_to_revised_target",
    "schedule_status",
    "overdue_days",
    "extension_count",
    "cost_overrun_delta_1m",
    "cost_overrun_delta_3m",
    "cost_overrun_trend_slope",
    "expenditure_ratio_delta_1m",
    "expenditure_ratio_delta_3m",
    "schedule_extension_delta_1m",
    "risk_signal_count",
    "page",
    "source_report",
    "data_quality_flag"
]


class MasterManager:
    """Manages reading, merging, validating, and saving the cumulative PAIMANA Master Excel dataset."""

    def __init__(self, master_path: str = DEFAULT_MASTER_PATH, config_path: str = "config/feature_config.yaml"):
        self.master_path = master_path
        self.config_path = config_path
        self.columns = self._load_schema()

    def _load_schema(self) -> List[str]:
        if os.path.exists(self.config_path):
            try:
                with open(self.config_path, "r", encoding="utf-8") as f:
                    cfg = yaml.safe_load(f)
                    return cfg.get("schema_columns", EXACT_55_COLUMNS)
            except Exception as e:
                logger.warning(f"Could not load schema from config: {e}. Using default 55 columns.")
        return EXACT_55_COLUMNS

    def load_existing_records(self) -> List[Dict[str, Any]]:
        """Load all historical records from the Master Excel file if it exists."""
        if not os.path.exists(self.master_path):
            return []

        try:
            df = pd.read_excel(self.master_path, sheet_name=SHEET_NAME, engine="openpyxl")
            # Replace NaN with None
            df = df.where(pd.notnull(df), None)
            records = df.to_dict(orient="records")
            logger.info(f"Loaded {len(records)} existing records from Master Excel.")
            return records
        except Exception as e:
            logger.error(f"Error loading existing master file {self.master_path}: {e}")
            return []

    def get_existing_keys(self) -> Set[Tuple[str, str]]:
        """Get set of (project_id, report_month) tuples already in master file."""
        records = self.load_existing_records()
        keys = set()
        for r in records:
            p_id = str(r.get("project_id", "")).strip()
            r_mo = str(r.get("report_month", "")).strip()
            if p_id and r_mo:
                keys.add((p_id, r_mo))
        return keys

    def merge_and_save(self, new_records: List[Dict[str, Any]], allow_reprocessing: bool = False, fresh: bool = False) -> Tuple[int, int, List[str]]:
        """
        Merge new records into master dataset and save to Excel.
        Returns (added_count, skipped_count, change_logs).
        """
        os.makedirs(os.path.dirname(self.master_path), exist_ok=True)
        existing_records = [] if fresh else self.load_existing_records()

        # Map by (project_id, report_month)
        record_map: Dict[Tuple[str, str], Dict[str, Any]] = {}
        for r in existing_records:
            p_id = str(r.get("project_id", "")).strip()
            r_mo = str(r.get("report_month", "")).strip()
            if p_id and r_mo:
                record_map[(p_id, r_mo)] = r

        added_count = 0
        skipped_count = 0
        reprocessed_count = 0
        change_logs = []

        for new_rec in new_records:
            p_id = str(new_rec.get("project_id", "")).strip()
            r_mo = str(new_rec.get("report_month", "")).strip()
            key = (p_id, r_mo)

            if key in record_map:
                if allow_reprocessing:
                    old_rec = record_map[key]
                    # Check differences
                    diffs = []
                    for col in self.columns:
                        if old_rec.get(col) != new_rec.get(col):
                            diffs.append(f"{col}: old='{old_rec.get(col)}' -> new='{new_rec.get(col)}'")
                    if diffs:
                        log_msg = f"Reprocessed project {p_id} for {r_mo}: {'; '.join(diffs)}"
                        change_logs.append(log_msg)
                        logger.info(log_msg)
                    record_map[key] = new_rec
                    reprocessed_count += 1
                else:
                    skipped_count += 1
                    logger.debug(f"Skipping duplicate record for ({p_id}, {r_mo}).")
            else:
                record_map[key] = new_rec
                added_count += 1

        # Combine all records and sort in ascending order by project_id and report_month
        def record_sort_key(r: Dict[str, Any]) -> Tuple[int, Any, str]:
            pid = str(r.get("project_id", "")).strip()
            r_mo = str(r.get("report_month", "")).strip()
            if pid.isdigit():
                return (0, int(pid), r_mo)
            return (1, pid, r_mo)

        all_records = list(record_map.values())
        all_records.sort(key=record_sort_key)

        self._save_to_excel(all_records)
        logger.info(f"Master file saved: {len(all_records)} total rows (+{added_count} added, {skipped_count} duplicates skipped, {reprocessed_count} reprocessed).")
        return added_count, skipped_count, change_logs

    def _save_to_excel(self, records: List[Dict[str, Any]]):
        """Write records to openpyxl workbook with professional formatting."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME

        # Write header
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for col_idx, col_name in enumerate(self.columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Write data rows
        data_font = Font(name="Calibri", size=10)
        data_align_left = Alignment(horizontal="left", vertical="center")
        data_align_right = Alignment(horizontal="right", vertical="center")
        data_align_center = Alignment(horizontal="center", vertical="center")

        date_columns = {"report_month", "approval_start", "original_target_doc", "revised_doc"}
        numeric_columns = {
            "project_age_months", "original_duration_months", "planned_remaining_months",
            "revised_remaining_months", "schedule_extension_months", "extension_rate_pct",
            "original_cost_crore", "revised_cost_crore", "cumulative_expenditure_crore",
            "cost_overrun_pct", "expenditure_ratio_pct", "cost_escalation_crore",
            "cost_escalation_ratio", "remaining_budget_crore", "expenditure_velocity_crore_month",
            "physical_progress_pct", "remaining_work_pct", "physical_progress_delta_1m",
            "physical_progress_delta_3m", "physical_progress_delta_6m", "progress_velocity_3m",
            "progress_velocity_6m", "progress_trend_slope", "physical_financial_gap_pct",
            "physical_to_expenditure_ratio", "progress_expenditure_mismatch_flag",
            "high_expenditure_low_progress_flag", "days_to_original_target",
            "days_to_revised_target", "overdue_days", "extension_count", "cost_overrun_delta_1m",
            "cost_overrun_delta_3m", "cost_overrun_trend_slope", "expenditure_ratio_delta_1m",
            "expenditure_ratio_delta_3m", "schedule_extension_delta_1m", "risk_signal_count", "page"
        }

        for row_idx, rec in enumerate(records, 2):
            for col_idx, col_name in enumerate(self.columns, 1):
                val = rec.get(col_name)
                # Format dates cleanly
                if col_name in date_columns and val is not None:
                    if hasattr(val, "strftime"):
                        val = val.strftime('%Y-%m-%d')
                    elif hasattr(val, "date") and callable(getattr(val, "date")):
                        val = val.date().strftime('%Y-%m-%d')
                    else:
                        val = str(val)[:10]

                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border

                if col_name in date_columns:
                    cell.alignment = data_align_center
                elif col_name in numeric_columns:
                    cell.alignment = data_align_right
                else:
                    cell.alignment = data_align_left

        # Freeze header row
        ws.freeze_panes = "A2"

        # Enable AutoFilter
        last_col_letter = get_column_letter(len(self.columns))
        last_row = max(len(records) + 1, 2)
        ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

        try:
            wb.save(self.master_path)
            logger.info(f"Saved Excel master to {self.master_path}")
        except PermissionError as e:
            backup_path = self.master_path.replace(".xlsx", "_latest.xlsx")
            logger.warning(f"Master file {self.master_path} is currently locked by another program (e.g. Microsoft Excel). Saving to {backup_path} instead.")
            wb.save(backup_path)
